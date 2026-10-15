import sys
from pathlib import Path
import openpyxl

# Allow importing when running directly as script or as module
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from backend.database import get_db_connection, init_db

BASE_DIR = Path(__file__).resolve().parent.parent
SEED_FILE = BASE_DIR / "data" / "events_seed.xlsx"


def seed_events():
    """Seed the events table from data/events_seed.xlsx."""
    # Ensure tables exist
    init_db()

    conn = get_db_connection()
    cursor = conn.cursor()

    # Skip if table already has rows
    cursor.execute("SELECT COUNT(*) FROM events")
    count = cursor.fetchone()[0]
    if count > 0:
        print(f"Events table already has {count} rows. Skipping seed.")
        conn.close()
        return

    # Check if seed file exists
    if not SEED_FILE.exists():
        raise FileNotFoundError(
            f"Seed data file not found at '{SEED_FILE}'. "
            "Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    wb = openpyxl.load_workbook(SEED_FILE)
    if "Events" not in wb.sheetnames:
        raise ValueError(f"Required sheet 'Events' was not found in '{SEED_FILE}'.")

    sheet = wb["Events"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        raise ValueError(f"No event rows found in '{SEED_FILE}'.")

    header = rows[0]
    print(f"Found header in spreadsheet: {header}")

    data_rows = rows[1:]
    inserted_count = 0

    for row in data_rows:
        if not row or all(v is None for v in row):
            continue

        event_id = row[0]
        title = row[1]
        category = row[2]
        description = row[3]
        event_date = row[4]
        event_time = row[5]
        venue = row[6]
        capacity = row[7]
        fee = row[8]
        icon = row[9]

        cursor.execute(
            """
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(event_id) if event_id is not None else None,
                str(title) if title is not None else "",
                str(category) if category is not None else "",
                str(description) if description is not None else "",
                str(event_date) if event_date is not None else "",
                str(event_time) if event_time is not None else "",
                str(venue) if venue is not None else "",
                int(capacity) if capacity is not None else 0,
                int(fee) if fee is not None else 0,
                str(icon) if icon is not None else "⚡"
            )
        )
        inserted_count += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {inserted_count} events from {SEED_FILE}.")


if __name__ == "__main__":
    seed_events()
    # Query and display seeded events
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events ORDER BY id ASC")
    rows = cursor.fetchall()
    print("\nCurrent contents of 'events' table:")
    print("-" * 90)
    print(f"{'ID':<4} | {'Title':<24} | {'Category':<12} | {'Date':<10} | {'Time':<10} | {'Cap':<4} | {'Fee':<4} | {'Icon'}")
    print("-" * 90)
    for r in rows:
        print(f"{r['id']:<4} | {r['title']:<24} | {r['category']:<12} | {r['event_date']:<10} | {r['event_time']:<10} | {r['capacity']:<4} | {r['fee']:<4} | {r['icon']}")
    print("-" * 90)
    print(f"Total rows in events table: {len(rows)}")
    conn.close()
